"""
migration_sgtg.py — Script de migration des données Excel SGTG vers PaieGabon SaaS
Usage : python migration_sgtg.py

Ce script :
1. Lit le fichier PAIE_SOCIETE_SGTG_2026.xlsx
2. Crée le tenant SGTG dans la base de données
3. Importe les 33 salariés
4. Importe les 218 bulletins de paie historiques
"""

import os, sys
from datetime import datetime, date
from openpyxl import load_workbook

# ── CONFIGURATION ──────────────────────────────────────────────────────────────
EXCEL_FILE   = "PAIE_SOCIETE_SGTG_2026.xlsx"   # Mettre le fichier dans le même dossier
ADMIN_EMAIL  = "admin@sgtg.ga"                  # Email admin SGTG
ADMIN_PASSWORD = "Sgtg2026!"                    # Mot de passe admin SGTG

# ── INITIALISATION FLASK ───────────────────────────────────────────────────────
from app import app
from models import db, Plan, Tenant, Utilisateur, CategorieEmploi, \
                   Salarie, Contrat, PeriodePaie, BulletinPaie
import secrets

def n(val, default=0):
    """Convertit une valeur en float, retourne default si None."""
    try:
        if val is None: return default
        return float(val)
    except: return default

def d(val):
    """Convertit une valeur en date."""
    if val is None: return None
    if isinstance(val, datetime): return val.date()
    if isinstance(val, date): return val
    try: return datetime.strptime(str(val)[:10], "%Y-%m-%d").date()
    except: return None

def migrer():
    print("=" * 60)
    print("  MIGRATION DONNÉES SGTG → PAIEGALON SAAS")
    print("=" * 60)

    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Fichier '{EXCEL_FILE}' introuvable !")
        print("   Placez le fichier Excel dans le même dossier que ce script.")
        sys.exit(1)

    print(f"\n📂 Lecture du fichier : {EXCEL_FILE}")
    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)

    with app.app_context():

        # ── ÉTAPE 1 : CRÉER LE TENANT SGTG ──────────────────────────────────
        print("\n🏢 Étape 1 — Création du tenant SGTG...")

        existing = Tenant.query.filter_by(slug="sgtg").first()
        if existing:
            print(f"   ℹ️  Tenant SGTG existe déjà (id={existing.id}) — mise à jour...")
            tenant = existing
        else:
            plan = Plan.query.filter_by(code="PRO").first()
            tenant = Tenant(slug="sgtg", statut="ACTIF", plan_id=plan.id if plan else None)
            db.session.add(tenant)

        # Infos société depuis Excel
        ws_soc = wb["INFOS SOCIETE"]
        infos = {}
        for row in ws_soc.iter_rows(values_only=True):
            if row[1] and row[2]:
                infos[str(row[1]).strip().upper()] = row[2]

        tenant.denomination  = str(infos.get("DENOMINATION SOCIALE", "SOCIETE SGTG")).strip().upper()
        tenant.sigle         = str(infos.get("SIGLE", "SGTG")).strip()
        tenant.activite      = str(infos.get("ACTIVITE", "GENIE CIVIL")).strip()
        tenant.secteur       = str(infos.get("SECTEUR", "")).strip()
        tenant.nif           = str(infos.get("NIF", "")).strip()
        tenant.adresse       = str(infos.get("ADRESSE", "")).strip()
        tenant.boite_postale = str(infos.get("BOÎTE POSTALE", infos.get("Boîte Postale", ""))).strip()
        tenant.ville         = str(infos.get("VILLE", "Libreville")).strip()
        tenant.region        = str(infos.get("REGION", "ESTUAIRE")).strip()
        tenant.pays          = str(infos.get("PAYS", "Gabon")).strip()
        if not tenant.token_api:
            tenant.token_api = secrets.token_hex(32)

        db.session.flush()
        print(f"   ✅ Tenant : {tenant.denomination} (id={tenant.id})")

        # ── ÉTAPE 2 : CRÉER L'ADMIN SGTG ────────────────────────────────────
        print("\n👤 Étape 2 — Création du compte admin SGTG...")
        admin_exist = Utilisateur.query.filter_by(email=ADMIN_EMAIL).first()
        if not admin_exist:
            admin = Utilisateur(
                nom="ADMINISTRATEUR", prenom="SGTG",
                email=ADMIN_EMAIL, role="TENANT_ADMIN",
                tenant_id=tenant.id, actif=True)
            admin.set_password(ADMIN_PASSWORD)
            db.session.add(admin)
            print(f"   ✅ Compte créé : {ADMIN_EMAIL} / {ADMIN_PASSWORD}")
        else:
            admin_exist.tenant_id = tenant.id
            print(f"   ℹ️  Compte existant mis à jour : {ADMIN_EMAIL}")

        # ── ÉTAPE 3 : CATÉGORIES D'EMPLOI ───────────────────────────────────
        print("\n📋 Étape 3 — Catégories d'emploi...")
        cats = {}
        for code, libelle in [
            ("C1", "Ouvriers / Conducteurs d'engins / Chauffeurs"),
            ("C2", "Techniciens"),
            ("C3", "Conducteurs de Travaux"),
            ("C4", "Cadres / Responsables / Gérants"),
        ]:
            cat = CategorieEmploi.query.filter_by(tenant_id=tenant.id, code=code).first()
            if not cat:
                cat = CategorieEmploi(tenant_id=tenant.id, code=code, libelle=libelle)
                db.session.add(cat)
                db.session.flush()
            cats[code] = cat
        print(f"   ✅ 4 catégories créées/vérifiées")

        # ── ÉTAPE 4 : IMPORT DES SALARIÉS ───────────────────────────────────
        print("\n👷 Étape 4 — Import des salariés...")
        ws_sal = wb["INFOS SALARIES"]
        header_sal = None
        nb_sal = 0
        salaries_map = {}  # matricule → objet Salarie

        for row in ws_sal.iter_rows(values_only=True):
            if not any(v is not None for v in row):
                continue
            if header_sal is None:
                header_sal = row
                continue
            if row[0] is None:
                continue

            matricule = str(row[0]).strip().upper()
            if not matricule:
                continue

            # Chercher si existe déjà
            sal = Salarie.query.filter_by(tenant_id=tenant.id, matricule=matricule).first()
            if not sal:
                sal = Salarie(tenant_id=tenant.id, matricule=matricule)
                db.session.add(sal)

            sal.nom                    = str(row[1]).strip().upper() if row[1] else "—"
            sal.prenom                 = str(row[2]).strip() if row[2] else "—"
            sal.telephone              = str(row[3]).strip() if row[3] else None
            sal.adresse                = str(row[4]).strip() if row[4] else None
            sal.nationalite            = str(row[5]).strip().upper() if row[5] else "GABONAISE"
            sal.sexe                   = str(row[6]).strip().upper() if row[6] else "M"
            sal.date_naissance         = d(row[7])
            sal.date_embauche          = d(row[9]) or date(2024, 8, 1)
            sal.date_cessation         = d(row[10])
            sal.situation_matrimoniale = str(row[11]).strip().upper() if row[11] else None
            sal.nb_enfants             = int(row[12]) if row[12] and str(row[12]).isdigit() else 0
            sal.nombre_parts           = float(row[13]) if row[13] else 1.0
            sal.numero_cnss            = str(row[14]).strip() if row[14] else None
            sal.numero_cnamgs          = str(row[15]).strip() if row[15] else None
            sal.emploi                 = str(row[16]).strip().upper() if row[16] else None
            sal.nb_enfants_moins_16ans = int(row[18]) if row[18] and str(row[18]).isdigit() else 0
            sal.assujetti_cnamgs       = str(row[19]).strip().upper() == "OUI" if row[19] else True
            sal.type_rupture           = str(row[20]).strip().upper() if row[20] else None
            sal.statut                 = "INACTIF" if sal.date_cessation else "ACTIF"

            # Catégorie
            cat_code = str(row[17]).strip().upper() if row[17] else "C1"
            sal.categorie_id = cats.get(cat_code, cats.get("C1")).id

            db.session.flush()
            salaries_map[matricule] = sal
            nb_sal += 1
            print(f"   {'✅' if sal.statut=='ACTIF' else '⚪'} {matricule} — {sal.nom} {sal.prenom} ({sal.emploi})")

        print(f"\n   📊 {nb_sal} salariés importés")

        # ── ÉTAPE 5 : IMPORT DES BULLETINS ──────────────────────────────────
        print("\n📄 Étape 5 — Import des bulletins de paie (218 lignes)...")
        ws_bul = wb["DONNEES DU BULLETIN"]
        MOIS_NOMS = {
            "JANVIER":1,"FÉVRIER":2,"MARS":3,"AVRIL":4,"MAI":5,"JUIN":6,
            "JUILLET":7,"AOÛT":8,"SEPTEMBRE":9,"OCTOBRE":10,"NOVEMBRE":11,"DÉCEMBRE":12,
            "AOUT":8,"FEVRIER":2
        }
        periodes_cache = {}
        nb_bul = 0
        nb_skip = 0

        for i, row in enumerate(ws_bul.iter_rows(values_only=True)):
            if i == 0: continue  # Entête
            if not any(v is not None for v in row): continue
            if row[4] is None: continue

            matricule = str(row[4]).strip().upper()
            if matricule not in salaries_map:
                nb_skip += 1
                continue

            annee = int(row[1]) if row[1] else None
            mois_str = str(row[2]).strip().upper() if row[2] else ""
            mois = MOIS_NOMS.get(mois_str)
            if not annee or not mois:
                nb_skip += 1
                continue

            # Période
            periode_key = f"{annee}-{mois}"
            if periode_key not in periodes_cache:
                p = PeriodePaie.query.filter_by(tenant_id=tenant.id, annee=annee, mois=mois).first()
                if not p:
                    mois_nom = [k for k,v in MOIS_NOMS.items() if v == mois and len(k) > 4][0]
                    trimestre = f"T{(mois-1)//3+1}"
                    p = PeriodePaie(tenant_id=tenant.id, annee=annee, mois=mois,
                                    libelle_mois=mois_nom, trimestre=trimestre, statut="CLÔTURÉ")
                    db.session.add(p)
                    db.session.flush()
                periodes_cache[periode_key] = p

            periode = periodes_cache[periode_key]
            sal = salaries_map[matricule]

            # Vérifier si bulletin existe
            bul = BulletinPaie.query.filter_by(
                tenant_id=tenant.id, salarie_id=sal.id, periode_id=periode.id).first()
            if not bul:
                bul = BulletinPaie(tenant_id=tenant.id, salarie_id=sal.id, periode_id=periode.id)
                db.session.add(bul)

            # Remplir les données
            bul.nb_jours_travailles   = int(n(row[42]))
            bul.salaire_base          = n(row[5])
            bul.heures_sup_10         = n(row[7])
            bul.heures_sup_30         = n(row[9])
            bul.heures_sup_40         = n(row[11])
            bul.heures_sup_70         = n(row[13])
            bul.absences              = n(row[15])
            bul.sursalaire            = n(row[17])
            bul.prime_caisse          = n(row[19])
            bul.carburant             = n(row[21])
            bul.prime_anciennete      = n(row[23])
            bul.indem_logement        = n(row[25])
            bul.indem_domesticite     = n(row[26])
            bul.indem_eau_electricite = n(row[27])
            bul.indem_nourriture      = n(row[28])
            bul.prime_rendement       = n(row[29])
            bul.prime_assiduité       = n(row[31])
            bul.prime_qualite         = n(row[33])
            bul.prime_performance     = n(row[35])
            bul.prime_transport       = n(row[37])
            bul.prime_responsabilite  = n(row[39])
            bul.allocations_conge     = n(row[41])
            bul.salaire_brut          = n(row[53])
            bul.base_cnss             = n(row[54])
            bul.cnss_salarie          = n(row[55])
            bul.cnss_patronale        = n(row[56])
            bul.base_cnamgs           = n(row[59])
            bul.cnamgs_salarie        = n(row[60])
            bul.cnamgs_patronale      = n(row[61])
            bul.fnh                   = n(row[62])
            bul.cfp                   = n(row[63])
            bul.base_tcs              = n(row[72])
            bul.tcs                   = n(row[73])
            bul.net_avant_irpp        = n(row[74])
            bul.base_irpp             = n(row[75])
            bul.irpp                  = n(row[76])
            bul.salaire_net           = n(row[77])
            bul.prime_panier          = n(row[78])
            bul.indem_transport       = n(row[79])
            bul.indem_representation  = n(row[80])
            bul.prime_salisure        = n(row[81])
            bul.acompte               = n(row[82])
            bul.net_a_payer           = n(row[83])
            bul.statut                = "VALIDÉ"
            bul.date_validation       = datetime.utcnow()

            nb_bul += 1

        # ── COMMIT FINAL ─────────────────────────────────────────────────────
        print(f"\n   📊 {nb_bul} bulletins importés ({nb_skip} ignorés)")
        print("\n💾 Enregistrement en base de données...")
        db.session.commit()

        # ── RÉSUMÉ ───────────────────────────────────────────────────────────
        print("\n" + "=" * 60)
        print("  ✅ MIGRATION TERMINÉE AVEC SUCCÈS !")
        print("=" * 60)
        print(f"\n  Tenant     : {tenant.denomination}")
        print(f"  Salariés   : {nb_sal}")
        print(f"  Bulletins  : {nb_bul}")
        print(f"  Périodes   : {len(periodes_cache)}")
        print(f"\n  Connexion  : {ADMIN_EMAIL}")
        print(f"  Mot de passe : {ADMIN_PASSWORD}")
        print(f"\n  URL : https://ameriack-paie.up.railway.app")
        print("=" * 60)

if __name__ == "__main__":
    migrer()
