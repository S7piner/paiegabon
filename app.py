"""
app.py — SaaS Paie Gabon — Multi-tenant
"""
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file, abort
from xhtml2pdf import pisa
import tempfile
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from datetime import datetime, date, timedelta
from functools import wraps
import io, os, secrets as sec

from models import (db, Plan, Tenant, Utilisateur, CategorieEmploi, Salarie,
                    Contrat, PeriodePaie, BulletinPaie, RubriquePaie, Conge,
                    Acompte, Journalier, Pointage, FeuillePaieJournalier)
from calculs_paie import calculer_bulletin, calculer_masse_salariale
from flask_mail import Mail, Message

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY","saas-paie-gabon-2026")
app.config["SQLALCHEMY_DATABASE_URI"] = os.environ.get("DATABASE_URL","sqlite:///saas_paie.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db.init_app(app)

# Configuration email Gmail
app.config["MAIL_SERVER"]   = "smtp.gmail.com"
app.config["MAIL_PORT"]     = 587
app.config["MAIL_USE_TLS"]  = True
app.config["MAIL_USERNAME"] = "wilsonzannou7@gmail.com"
app.config["MAIL_PASSWORD"] = "exvgzgcnyczsumwp"
app.config["MAIL_DEFAULT_SENDER"] = "wilsonzannou7@gmail.com"

app.config["MAIL_SUPPRESS_SEND"] = False
mail = Mail(app)

login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Veuillez vous connecter."

@login_manager.user_loader
def load_user(uid): return Utilisateur.query.get(int(uid))

def super_admin_required(f):
    @wraps(f)
    def d(*a,**k):
        if not current_user.is_authenticated or not current_user.is_super_admin: abort(403)
        return f(*a,**k)
    return d

def _parse_date(val):
    if not val: return None
    if isinstance(val, date): return val
    if isinstance(val, datetime): return val.date()
    try: return datetime.strptime(str(val).strip()[:10], "%Y-%m-%d").date()
    except: return None

def calculer_parts_irpp(situation_matrimoniale: str, nb_enfants: int) -> float:
    situation = (situation_matrimoniale or "").upper().strip()
    nb_enf = int(nb_enfants or 0)
    if "CELIBATAIRE" in situation and "AEAC" in situation: parts = 1.5
    elif "CELIBATAIRE" in situation: parts = 1.0
    elif "DIVORCE" in situation and "AEAC" in situation: parts = 1.5
    elif "DIVORCE" in situation: parts = 1.0
    elif "MARIE" in situation or "MARIÉ" in situation: parts = 2.0
    elif "VEUF" in situation and "2 ANS" in situation: parts = 1.5
    elif "VEUF" in situation: parts = 2.0
    else: parts = 1.0
    parts += nb_enf * 0.5
    return round(parts, 1)

def tenant_required(f):
    @wraps(f)
    def d(*a,**k):
        if not current_user.is_authenticated: return redirect(url_for("login"))
        if not current_user.is_super_admin and (not current_user.tenant_id or not current_user.tenant or current_user.tenant.statut not in ("ACTIF","ESSAI")):
            flash("Compte suspendu ou non associé à une entreprise.","error")
            return redirect(url_for("login"))
        return f(*a,**k)
    return d

def can_edit(f):
    @wraps(f)
    def d(*a,**k):
        if not current_user.can_edit: abort(403)
        return f(*a,**k)
    return d

def get_tenant():
    if current_user.is_super_admin: return None
    return current_user.tenant

# ── Auth ──────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("admin_dashboard") if current_user.is_super_admin else url_for("dashboard"))
    return redirect(url_for("login"))

@app.route("/login", methods=["GET","POST"])
def login():
    if current_user.is_authenticated: return redirect(url_for("index"))
    if request.method == "POST":
        email = request.form.get("email","").strip().lower()
        pw    = request.form.get("password","")
        user  = Utilisateur.query.filter_by(email=email, actif=True).first()
        if user and user.check_password(pw):
            login_user(user)
            user.derniere_connexion = datetime.utcnow()
            db.session.commit()
            return redirect(url_for("index"))
        flash("Email ou mot de passe incorrect.","error")
    return render_template("auth/login.html")

@app.route("/inscription", methods=["GET","POST"])
def inscription():
    plans = Plan.query.filter_by(actif=True).all()
    if request.method == "POST":
        email = request.form.get("email","").strip().lower()
        if Utilisateur.query.filter_by(email=email).first():
            flash("Email déjà utilisé.","error")
            return render_template("auth/inscription.html", plans=plans)
        plan = Plan.query.get(request.form.get("plan_id","")) or Plan.query.filter_by(code="STARTER").first()
        denom = request.form.get("denomination","").strip()
        slug_base = denom.lower().replace(" ","_")[:30]
        slug = slug_base; i=1
        while Tenant.query.filter_by(slug=slug).first(): slug=f"{slug_base}_{i}"; i+=1
        t = Tenant(slug=slug, denomination=denom.upper(),
                   sigle=request.form.get("sigle","").strip().upper(),
                   activite=request.form.get("activite","").strip(),
                   nif=request.form.get("nif","").strip(),
                   telephone=request.form.get("telephone","").strip(),
                   ville=request.form.get("ville","Libreville"),
                   pays="Gabon", plan_id=plan.id if plan else None,
                   statut="ESSAI", date_expiration=datetime.utcnow()+timedelta(days=30))
        t.token_api = sec.token_hex(32)
        db.session.add(t); db.session.flush()
        for code,lib in [("C1","Ouvriers"),("C2","Techniciens"),("C3","Conducteurs de Travaux"),("C4","Cadres")]:
            db.session.add(CategorieEmploi(tenant_id=t.id,code=code,libelle=lib))
        admin = Utilisateur(nom=request.form.get("nom","").strip().upper(),
                            prenom=request.form.get("prenom","").strip(),
                            email=email, role="TENANT_ADMIN", tenant_id=t.id, actif=True)
        admin.set_password(request.form.get("password",""))
        db.session.add(admin); db.session.commit()
        flash("Bienvenue ! Essai gratuit de 30 jours activé.","success")
        login_user(admin)
        return redirect(url_for("dashboard"))
    return render_template("auth/inscription.html", plans=plans)

@app.route("/logout")
@login_required
def logout():
    logout_user(); return redirect(url_for("login"))

# ── Super-Admin ───────────────────────────────────────────────────────────────
@app.route("/admin")
@super_admin_required
def admin_dashboard():
    tenants   = Tenant.query.order_by(Tenant.date_inscription.desc()).all()
    total_sal = db.session.query(db.func.count(Salarie.id)).scalar() or 0
    total_bul = db.session.query(db.func.count(BulletinPaie.id)).scalar() or 0
    revenus   = sum((float(t.plan.prix_mensuel) if t.plan else 0) for t in tenants if t.statut=="ACTIF")
    return render_template("admin/dashboard.html",
        tenants=tenants, nb_actifs=sum(1 for t in tenants if t.statut=="ACTIF"),
        nb_essai=sum(1 for t in tenants if t.statut=="ESSAI"),
        total_sal=total_sal, total_bul=total_bul, revenus=revenus)

@app.route("/admin/tenants")
@super_admin_required
def admin_tenants():
    q=request.args.get("q",""); statut=request.args.get("statut","")
    query=Tenant.query
    if q: query=query.filter(Tenant.denomination.ilike(f"%{q}%"))
    if statut: query=query.filter_by(statut=statut)
    return render_template("admin/tenants.html", tenants=query.order_by(Tenant.date_inscription.desc()).all(),
        plans=Plan.query.all(), q=q, statut=statut)

@app.route("/admin/tenants/<int:id>")
@super_admin_required
def admin_tenant_detail(id):
    t = Tenant.query.get_or_404(id)
    return render_template("admin/tenant_detail.html", tenant=t,
        nb_salaries=Salarie.query.filter_by(tenant_id=id).count(),
        nb_bulletins=BulletinPaie.query.filter_by(tenant_id=id).count(),
        users=Utilisateur.query.filter_by(tenant_id=id).all(),
        plans=Plan.query.all())

@app.route("/admin/tenants/<int:id>/statut", methods=["POST"])
@super_admin_required
def admin_tenant_statut(id):
    t=Tenant.query.get_or_404(id)
    t.statut=request.form.get("statut",t.statut)
    if request.form.get("plan_id"): t.plan_id=int(request.form["plan_id"])
    db.session.commit(); flash(f"{t.denomination} mis à jour.","success")
    return redirect(url_for("admin_tenant_detail",id=id))

@app.route("/admin/tenants/<int:id>/impersonate")
@super_admin_required
def admin_impersonate(id):
    u=Utilisateur.query.filter_by(tenant_id=id,role="TENANT_ADMIN").first()
    if not u: flash("Aucun admin trouvé.","error"); return redirect(url_for("admin_tenant_detail",id=id))
    logout_user(); login_user(u)
    flash(f"Connecté en tant que {u.nom_complet} ({u.tenant.denomination})","warning")
    return redirect(url_for("dashboard"))

@app.route("/admin/plans", methods=["GET","POST"])
@super_admin_required
def admin_plans():
    if request.method=="POST":
        p=Plan(code=request.form["code"].strip().upper(),
               nom=request.form["nom"].strip(),
               prix_mensuel=float(request.form.get("prix_mensuel",0)),
               max_salaries=int(request.form["max_salaries"]) if request.form.get("max_salaries") else None,
               max_utilisateurs=int(request.form["max_utilisateurs"]) if request.form.get("max_utilisateurs") else None,
               description=request.form.get("description",""))
        db.session.add(p); db.session.commit(); flash("Plan créé.","success")
    return render_template("admin/plans.html", plans=Plan.query.all())

@app.route("/admin/import", methods=["GET","POST"])
@login_required
@super_admin_required
def admin_import_excel():
    from werkzeug.utils import secure_filename
    tenants = Tenant.query.order_by(Tenant.denomination).all()
    resultats = None
    if request.method == "POST":
        tenant_id = request.form.get("tenant_id", type=int)
        tenant = Tenant.query.get_or_404(tenant_id)
        f = request.files.get("excel_file")
        if not f or not f.filename.endswith(".xlsx"):
            flash("Fichier invalide. Utilisez un fichier .xlsx", "error")
            return render_template("admin/import_excel.html", tenants=tenants)
        imp_societe  = "import_societe"  in request.form
        imp_salaries = "import_salaries" in request.form
        imp_bulletins= "import_bulletins"in request.form
        ecraser      = "ecraser"         in request.form
        try:
            from openpyxl import load_workbook
            from datetime import datetime as dt2, date as d2
            wb = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
            nb_sal=nb_bul=nb_per=nb_err=0
            def nv(v,d=0):
                try: return float(v) if v is not None else d
                except: return d
            def dv(v):
                if v is None: return None
                if isinstance(v,dt2): return v.date()
                if isinstance(v,d2): return v
                try: return dt2.strptime(str(v)[:10],"%Y-%m-%d").date()
                except: return None
            if imp_societe and "INFOS SOCIETE" in wb.sheetnames:
                ws=wb["INFOS SOCIETE"]; infos={}
                for row in ws.iter_rows(values_only=True):
                    if row[1] and row[2]: infos[str(row[1]).strip().upper()]=row[2]
                tenant.denomination=str(infos.get("DENOMINATION SOCIALE",tenant.denomination)).strip().upper()
                tenant.sigle=str(infos.get("SIGLE",tenant.sigle or "")).strip()
                tenant.activite=str(infos.get("ACTIVITE",tenant.activite or "")).strip()
                tenant.nif=str(infos.get("NIF",tenant.nif or "")).strip()
                tenant.adresse=str(infos.get("ADRESSE",tenant.adresse or "")).strip()
                tenant.ville=str(infos.get("VILLE",tenant.ville or "Libreville")).strip()
            cats={}
            for code,lib in [("C1","Ouvriers"),("C2","Techniciens"),("C3","Conducteurs"),("C4","Cadres")]:
                cat=CategorieEmploi.query.filter_by(tenant_id=tenant.id,code=code).first()
                if not cat:
                    cat=CategorieEmploi(tenant_id=tenant.id,code=code,libelle=lib)
                    db.session.add(cat); db.session.flush()
                cats[code]=cat
            salaries_map={}
            if imp_salaries and "INFOS SALARIES" in wb.sheetnames:
                ws=wb["INFOS SALARIES"]; header=None
                for row in ws.iter_rows(values_only=True):
                    if not any(v is not None for v in row): continue
                    if header is None: header=row; continue
                    if row[0] is None: continue
                    matricule=str(row[0]).strip().upper()
                    if not matricule: continue
                    sal=Salarie.query.filter_by(tenant_id=tenant.id,matricule=matricule).first()
                    if sal and not ecraser: salaries_map[matricule]=sal; continue
                    if not sal: sal=Salarie(tenant_id=tenant.id,matricule=matricule); db.session.add(sal)
                    sal.nom=str(row[1]).strip().upper() if row[1] else "—"
                    sal.prenom=str(row[2]).strip() if row[2] else "—"
                    sal.telephone=str(row[3]).strip() if row[3] else None
                    sal.nationalite=str(row[5]).strip().upper() if row[5] else "GABONAISE"
                    sal.sexe=str(row[6]).strip().upper() if row[6] else "M"
                    sal.date_naissance=dv(row[7])
                    sal.date_embauche=dv(row[9]) or d2(2024,8,1)
                    sal.date_cessation=dv(row[10])
                    sal.situation_matrimoniale=str(row[11]).strip().upper() if row[11] else None
                    sal.nb_enfants=int(row[12]) if row[12] and str(row[12]).replace('.','').isdigit() else 0
                    sal.nombre_parts=float(row[13]) if row[13] else 1.0
                    sal.numero_cnss=str(row[14]).strip() if row[14] else None
                    sal.numero_cnamgs=str(row[15]).strip() if row[15] else None
                    sal.emploi=str(row[16]).strip().upper() if row[16] else None
                    sal.nb_enfants_moins_16ans=int(row[18]) if row[18] and str(row[18]).replace('.','').isdigit() else 0
                    sal.assujetti_cnamgs=str(row[19]).strip().upper()=="OUI" if row[19] else True
                    sal.statut="INACTIF" if dv(row[10]) else "ACTIF"
                    cat_code=str(row[17]).strip().upper() if row[17] else "C1"
                    sal.categorie_id=cats.get(cat_code,cats.get("C1")).id
                    db.session.flush(); salaries_map[matricule]=sal; nb_sal+=1
            if not imp_salaries:
                for s in Salarie.query.filter_by(tenant_id=tenant.id).all():
                    salaries_map[s.matricule]=s
            MOIS={"JANVIER":1,"FÉVRIER":2,"FEVRIER":2,"MARS":3,"AVRIL":4,"MAI":5,"JUIN":6,
                  "JUILLET":7,"AOÛT":8,"AOUT":8,"SEPTEMBRE":9,"OCTOBRE":10,"NOVEMBRE":11,"DÉCEMBRE":12,"DECEMBRE":12}
            periodes_cache={}
            if imp_bulletins and "DONNEES DU BULLETIN" in wb.sheetnames:
                ws=wb["DONNEES DU BULLETIN"]
                for i,row in enumerate(ws.iter_rows(values_only=True)):
                    if i==0: continue
                    if not any(v is not None for v in row): continue
                    if row[4] is None: continue
                    matricule=str(row[4]).strip().upper().replace(' - ','-').replace('- ','-').replace(' -','-')
                    if matricule not in salaries_map: nb_err+=1; continue
                    annee=int(row[1]) if row[1] else None
                    mois_str=str(row[2]).strip().upper() if row[2] else ""
                    mois=MOIS.get(mois_str)
                    if not annee or not mois: nb_err+=1; continue
                    pk=f"{annee}-{mois}"
                    if pk not in periodes_cache:
                        p=PeriodePaie.query.filter_by(tenant_id=tenant.id,annee=annee,mois=mois).first()
                        if not p:
                            mois_nom=[k for k,v in MOIS.items() if v==mois and len(k)>4][0]
                            p=PeriodePaie(tenant_id=tenant.id,annee=annee,mois=mois,
                                libelle_mois=mois_nom,trimestre=f"T{(mois-1)//3+1}",statut="CLÔTURÉ")
                            db.session.add(p); db.session.flush(); nb_per+=1
                        periodes_cache[pk]=p
                    sal=salaries_map[matricule]
                    bul=BulletinPaie.query.filter_by(tenant_id=tenant.id,salarie_id=sal.id,periode_id=periodes_cache[pk].id).first()
                    if bul and not ecraser: continue
                    if not bul: bul=BulletinPaie(tenant_id=tenant.id,salarie_id=sal.id,periode_id=periodes_cache[pk].id); db.session.add(bul)
                    bul.nb_jours_travailles=int(nv(row[42]))
                    bul.salaire_base=nv(row[5]); bul.heures_sup_10=nv(row[7]); bul.heures_sup_30=nv(row[9])
                    bul.heures_sup_40=nv(row[11]); bul.heures_sup_70=nv(row[13]); bul.absences=nv(row[15])
                    bul.sursalaire=nv(row[17]); bul.prime_caisse=nv(row[19]); bul.carburant=nv(row[21])
                    bul.prime_anciennete=nv(row[23]); bul.indem_logement=nv(row[25])
                    bul.indem_domesticite=nv(row[26]); bul.indem_eau_electricite=nv(row[27])
                    bul.indem_nourriture=nv(row[28]); bul.prime_rendement=nv(row[29])
                    bul.prime_assiduité=nv(row[31]); bul.prime_qualite=nv(row[33])
                    bul.prime_performance=nv(row[35]); bul.prime_transport=nv(row[37])
                    bul.prime_responsabilite=nv(row[39]); bul.allocations_conge=nv(row[41])
                    bul.salaire_brut=nv(row[53]); bul.base_cnss=nv(row[54])
                    bul.cnss_salarie=nv(row[55]); bul.cnss_patronale=nv(row[56])
                    bul.base_cnamgs=nv(row[59]); bul.cnamgs_salarie=nv(row[60]); bul.cnamgs_patronale=nv(row[61])
                    bul.fnh=nv(row[62]); bul.cfp=nv(row[63]); bul.base_tcs=nv(row[72]); bul.tcs=nv(row[73])
                    bul.net_avant_irpp=nv(row[74]); bul.base_irpp=nv(row[75]); bul.irpp=nv(row[76])
                    bul.salaire_net=nv(row[77]); bul.prime_panier=nv(row[78]); bul.indem_transport=nv(row[79])
                    bul.indem_representation=nv(row[80]); bul.prime_salisure=nv(row[81])
                    bul.acompte=nv(row[82]); bul.net_a_payer=nv(row[83]) if len(row)>83 else 0
                    bul.statut="VALIDÉ"; bul.date_validation=datetime.utcnow(); nb_bul+=1
            db.session.commit()
            resultats={"nb_salaries":nb_sal,"nb_bulletins":nb_bul,"nb_periodes":nb_per,"erreurs":nb_err}
            flash(f"Import réussi ! {nb_sal} salariés, {nb_bul} bulletins, {nb_per} périodes.","success")
        except Exception as e:
            db.session.rollback()
            flash(f"Erreur : {str(e)}","error")
    return render_template("admin/import_excel.html", tenants=tenants, resultats=resultats)

@app.route("/admin/update-taux", methods=["POST"])
@super_admin_required
def admin_update_taux():
    taux = {"CNSS":(0.05,0.18),"CNAMGS":(0.02,0.041),"FNH":(0.0,0.03),"TCS":(0.05,0.0),"CFP":(0.0,0.005)}
    nb = 0
    for code,(sal,pat) in taux.items():
        for r in RubriquePaie.query.filter_by(code=code).all():
            r.taux_salarie=sal; r.taux_patronal=pat; nb+=1
    db.session.commit()
    flash(f"Taux mis a jour ({nb} rubriques).", "success")
    return redirect(url_for("admin_rubriques"))

@app.route("/admin/rubriques", methods=["GET","POST"])
@super_admin_required
def admin_rubriques():
    if request.method=="POST":
        r=RubriquePaie(code=request.form["code"].strip().upper(),
            libelle=request.form["libelle"].strip(), type=request.form.get("type","COTISATION"),
            taux_salarie=float(request.form["taux_salarie"]) if request.form.get("taux_salarie") else None,
            taux_patronal=float(request.form["taux_patronal"]) if request.form.get("taux_patronal") else None,
            plafond_mensuel=float(request.form["plafond_mensuel"]) if request.form.get("plafond_mensuel") else None)
        db.session.add(r); db.session.commit(); flash("Rubrique créée.","success")
    return render_template("admin/rubriques.html", rubriques=RubriquePaie.query.all())

@app.route("/admin/tenants/<int:id>/supprimer", methods=["POST"])
@super_admin_required
def admin_tenant_supprimer(id):
    t = Tenant.query.get_or_404(id)
    nom = t.denomination
    try:
        for s in Salarie.query.filter_by(tenant_id=id).all():
            BulletinPaie.query.filter_by(salarie_id=s.id).delete()
            Contrat.query.filter_by(salarie_id=s.id).delete()
            Pointage.query.filter_by(salarie_id=s.id).delete()
            Acompte.query.filter_by(salarie_id=s.id).delete()
            Conge.query.filter_by(salarie_id=s.id).delete()
        Salarie.query.filter_by(tenant_id=id).delete()
        PeriodePaie.query.filter_by(tenant_id=id).delete()
        CategorieEmploi.query.filter_by(tenant_id=id).delete()
        Utilisateur.query.filter_by(tenant_id=id).delete()
        Journalier.query.filter_by(tenant_id=id).delete()
        Acompte.query.filter_by(tenant_id=id).delete()
        Conge.query.filter_by(tenant_id=id).delete()
        db.session.delete(t); db.session.commit()
        flash(f"Entreprise {nom} supprimée.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Erreur: {str(e)}", "error")
    return redirect(url_for("admin_tenants"))

# ── Dashboard ─────────────────────────────────────────────────────────────────
@app.route("/dashboard")
@login_required
def dashboard():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: flash("Aucune entreprise associée.","error"); return redirect(url_for("login"))
    now=datetime.now()
    nb_actifs   = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").count()
    nb_inactifs = Salarie.query.filter_by(tenant_id=t.id, statut="INACTIF").count()
    nb_total    = Salarie.query.filter_by(tenant_id=t.id).count()
    debut_mois  = datetime(now.year, now.month, 1).date()
    nb_new_mois = Salarie.query.filter(Salarie.tenant_id==t.id, Salarie.date_embauche>=debut_mois).count()
    periode = PeriodePaie.query.filter_by(tenant_id=t.id, annee=now.year, mois=now.month).first()
    masse={}; nb_v=nb_b=nb_p=0
    if periode:
        buls = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=periode.id).all()
        masse = calculer_masse_salariale(buls)
        nb_v = sum(1 for b in buls if b.statut=="VALIDÉ")
        nb_p = sum(1 for b in buls if b.statut=="PAYÉ")
        nb_b = sum(1 for b in buls if b.statut=="BROUILLON")
    evolution = []
    for i in range(5, -1, -1):
        m = now.month - i; y = now.year
        while m <= 0: m += 12; y -= 1
        p = PeriodePaie.query.filter_by(tenant_id=t.id, annee=y, mois=m).first()
        mois_noms = ["","Jan","Fév","Mar","Avr","Mai","Jun","Jul","Aoû","Sep","Oct","Nov","Déc"]
        if p:
            buls_p = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=p.id).all()
            total_net = sum(float(b.net_a_payer or 0) for b in buls_p)
            total_brut = sum(float(b.salaire_brut or 0) for b in buls_p)
            total_charges = sum(float(b.cnss_patronale or 0)+float(b.cnamgs_patronale or 0)+float(b.fnh or 0)+float(b.cfp or 0) for b in buls_p)
        else:
            total_net=total_brut=total_charges=0; buls_p=[]
        evolution.append({"mois":mois_noms[m],"annee":y,"brut":round(total_brut),"net":round(total_net),"charges":round(total_charges),"nb_bulletins":len(buls_p)})
    top_salaries = []
    if periode:
        top_salaries = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=periode.id).order_by(BulletinPaie.net_a_payer.desc()).limit(5).all()
    from sqlalchemy import func
    cats_stats = db.session.query(CategorieEmploi.code, CategorieEmploi.libelle, func.count(Salarie.id).label("nb"))\
        .join(Salarie, Salarie.categorie_id==CategorieEmploi.id)\
        .filter(Salarie.tenant_id==t.id, Salarie.statut=="ACTIF")\
        .group_by(CategorieEmploi.code, CategorieEmploi.libelle).all()
    derniers = BulletinPaie.query.filter_by(tenant_id=t.id).order_by(BulletinPaie.date_creation.desc()).limit(6).all()
    alertes = []
    if nb_b > 0: alertes.append({"type":"warning","msg":f"{nb_b} bulletin(s) en brouillon à valider"})
    if not periode: alertes.append({"type":"info","msg":f"Aucune période ouverte pour {PeriodePaie.MOIS_NOMS[now.month]} {now.year}"})
    return render_template("tenant/dashboard.html", tenant=t,
        nb_actifs=nb_actifs, nb_inactifs=nb_inactifs, nb_total=nb_total,
        nb_new_mois=nb_new_mois, periode=periode, masse=masse,
        nb_valides=nb_v, nb_payes=nb_p, nb_brouillon=nb_b,
        evolution=evolution, top_salaries=top_salaries,
        cats_stats=cats_stats, derniers=derniers, alertes=alertes, now=now)

# ── Salariés ──────────────────────────────────────────────────────────────────
@app.route("/salaries")
@login_required
def salaries():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    q=request.args.get("q",""); statut=request.args.get("statut","")
    query=Salarie.query.filter_by(tenant_id=t.id)
    if q: query=query.filter(db.or_(Salarie.nom.ilike(f"%{q}%"),Salarie.prenom.ilike(f"%{q}%"),Salarie.matricule.ilike(f"%{q}%")))
    if statut: query=query.filter_by(statut=statut)
    return render_template("tenant/salaries.html", salaries=query.order_by(Salarie.nom).all(),
        categories=CategorieEmploi.query.filter_by(tenant_id=t.id).all(), q=q, statut=statut, tenant=t)

@app.route("/salaries/nouveau", methods=["GET","POST"])
@login_required
def salarie_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.can_edit: abort(403)
    cats=CategorieEmploi.query.filter_by(tenant_id=t.id).all()
    if request.method=="POST":
        if not t.est_dans_limite:
            flash(f"Limite atteinte ({t.plan.max_salaries} salariés). Passez au plan supérieur.","error")
            return redirect(url_for("salaries"))
        s=Salarie(tenant_id=t.id,
            matricule=request.form["matricule"].strip().upper(),
            categorie_id=request.form.get("categorie_id") or None,
            nom=request.form["nom"].strip().upper(), prenom=request.form["prenom"].strip(),
            telephone=request.form.get("telephone"), email=request.form.get("email","").strip() or None,
            nationalite=request.form.get("nationalite","GABONAISE"),
            sexe=request.form.get("sexe"),
            date_naissance=_pd(request.form.get("date_naissance")),
            date_embauche=_pd(request.form["date_embauche"]),
            situation_matrimoniale=request.form.get("situation_matrimoniale"),
            nb_enfants=int(request.form.get("nb_enfants") or 0),
            nb_enfants_moins_16ans=int(request.form.get("nb_enfants_moins_16ans") or 0),
            nombre_parts=float(request.form.get("nombre_parts") or 1),
            numero_cnss=request.form.get("numero_cnss"), numero_cnamgs=request.form.get("numero_cnamgs"),
            emploi=request.form.get("emploi"), assujetti_cnamgs=request.form.get("assujetti_cnamgs")=="OUI", statut="ACTIF")
        db.session.add(s)
        sb=float(request.form.get("salaire_base") or 0)
        if sb: db.session.add(Contrat(tenant_id=t.id,salarie=s,type_contrat=request.form.get("type_contrat","CDI"),date_debut=s.date_embauche,salaire_base=sb,poste=s.emploi,actif=True))
        db.session.commit(); flash(f"Salarié {s.nom_complet} créé.","success")
        return redirect(url_for("salarie_detail",id=s.id))
    return render_template("tenant/salarie_form.html", salarie=None, categories=cats, action="nouveau", tenant=t)

@app.route("/salaries/<int:id>")
@login_required
def salarie_detail(id):
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    s = Salarie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    bulletins = BulletinPaie.query.filter_by(salarie_id=id, tenant_id=t.id).order_by(BulletinPaie.date_creation.desc()).all()
    contrat = Contrat.query.filter_by(salarie_id=id, tenant_id=t.id, actif=True).first()
    conge = Conge.query.filter_by(salarie_id=id, tenant_id=t.id, annee=datetime.now().year).first()
    total_brut = sum(float(b.salaire_brut or 0) for b in bulletins)
    total_net  = sum(float(b.net_a_payer or 0) for b in bulletins)
    total_cnss = sum(float(b.cnss_salarie or 0) for b in bulletins)
    total_irpp = sum(float(b.irpp or 0) for b in bulletins)
    nb_mois    = len(bulletins)
    anciennete_jours = (datetime.now().date() - s.date_embauche).days if s.date_embauche else 0
    anciennete_ans   = anciennete_jours // 365
    anciennete_mois  = (anciennete_jours % 365) // 30
    return render_template("tenant/salarie_detail.html",
        salarie=s, tenant=t, bulletins=bulletins, contrat=contrat, conge=conge,
        total_brut=total_brut, total_net=total_net, total_cnss=total_cnss,
        total_irpp=total_irpp, nb_mois=nb_mois,
        anciennete_ans=anciennete_ans, anciennete_mois=anciennete_mois)

@app.route("/salaries/<int:id>/modifier", methods=["GET","POST"])
@login_required
def salarie_modifier(id):
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.can_edit: abort(403)
    s = Salarie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    cats = CategorieEmploi.query.filter_by(tenant_id=t.id).all()
    if request.method=="POST":
        for f,v in [("nom",request.form["nom"].strip().upper()),("prenom",request.form["prenom"].strip()),
            ("telephone",request.form.get("telephone")),
            ("email",request.form.get("email","").strip() or None),
            ("nationalite",request.form.get("nationalite")),
            ("sexe",request.form.get("sexe")),("date_naissance",_pd(request.form.get("date_naissance"))),
            ("situation_matrimoniale",request.form.get("situation_matrimoniale")),
            ("nb_enfants",int(request.form.get("nb_enfants") or 0)),
            ("nombre_parts",calculer_parts_irpp(request.form.get("situation_matrimoniale",""),int(request.form.get("nb_enfants",0) or 0))),
            ("numero_cnss",request.form.get("numero_cnss")),("numero_cnamgs",request.form.get("numero_cnamgs")),
            ("emploi",request.form.get("emploi")),("categorie_id",request.form.get("categorie_id") or None),
            ("statut",request.form.get("statut","ACTIF")),("date_modification",datetime.utcnow())]:
            setattr(s,f,v)
        db.session.commit(); flash("Fiche mise à jour.","success")
        return redirect(url_for("salarie_detail",id=s.id))
    return render_template("tenant/salarie_form.html", salarie=s, categories=cats, action="modifier", tenant=t)

# ── Bulletins ─────────────────────────────────────────────────────────────────
@app.route("/bulletins")
@login_required
def bulletins():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    pid=request.args.get("periode_id",type=int); sf=request.args.get("statut","")
    periodes=PeriodePaie.query.filter_by(tenant_id=t.id).order_by(PeriodePaie.annee.desc(),PeriodePaie.mois.desc()).all()
    ps=None; buls=[]; masse={}
    if pid:
        ps=PeriodePaie.query.filter_by(id=pid,tenant_id=t.id).first_or_404()
        q=BulletinPaie.query.filter_by(periode_id=pid,tenant_id=t.id)
        if sf: q=q.filter_by(statut=sf)
        buls=q.join(Salarie).order_by(Salarie.nom).all()
        masse=calculer_masse_salariale(buls)
    return render_template("tenant/bulletins.html", periodes=periodes, periode_sel=ps,
        bulletins=buls, masse=masse, statut_filtre=sf, tenant=t)

@app.route("/bulletins/saisie", methods=["GET","POST"])
@login_required
def bulletin_saisie():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.can_edit: abort(403)
    sals=Salarie.query.filter_by(tenant_id=t.id,statut="ACTIF").order_by(Salarie.nom).all()
    pers=PeriodePaie.query.filter_by(tenant_id=t.id,statut="OUVERT").order_by(PeriodePaie.annee.desc(),PeriodePaie.mois.desc()).all()
    if request.method=="POST":
        sid=int(request.form["salarie_id"]); pid=int(request.form["periode_id"])
        s=Salarie.query.filter_by(id=sid,tenant_id=t.id).first_or_404()
        periode = PeriodePaie.query.filter_by(id=pid, tenant_id=t.id).first_or_404()
        # Auto-déduire les acomptes EN_ATTENTE du salarié pour ce mois
        acomptes_en_attente = Acompte.query.filter_by(
            tenant_id=t.id, salarie_id=sid,
            mois=periode.mois, annee=periode.annee, statut="EN_ATTENTE").all()
        total_acomptes = sum(float(a.montant) for a in acomptes_en_attente)
        donnees={k:float(v) if v else 0 for k,v in request.form.items() if k not in("salarie_id","periode_id","csrf_token","action","nb_jours_travailles")}
        # Injecter le total des acomptes
        if total_acomptes > 0:
            donnees["acompte"] = max(donnees.get("acompte", 0), total_acomptes)
        res=calculer_bulletin(donnees,nb_parts=float(s.nombre_parts or 1))
        ex=BulletinPaie.query.filter_by(tenant_id=t.id,salarie_id=sid,periode_id=pid).first()
        b=ex or BulletinPaie(tenant_id=t.id,salarie_id=sid,periode_id=pid)
        if not ex: db.session.add(b)
        for k,v in res.items():
            if not k.startswith("_") and hasattr(b,k): setattr(b,k,v)
        b.nb_jours_travailles=int(request.form.get("nb_jours_travailles") or 0)
        action=request.form.get("action","brouillon")
        if action=="valider":
            b.statut="VALIDÉ"; b.date_validation=datetime.utcnow()
            # Marquer les acomptes comme DÉDUITS
            for a in acomptes_en_attente:
                a.statut = "DEDUIT"
        else:
            b.statut="BROUILLON"
        db.session.commit()
        if total_acomptes > 0:
            flash(f"Bulletin sauvegardé. Acompte de {int(total_acomptes):,} FCFA déduit automatiquement.".replace(",", " "), "success")
        else:
            flash(f"Bulletin {'validé' if b.statut=='VALIDÉ' else 'sauvegardé'}.","success")
        return redirect(url_for("bulletin_detail",id=b.id))
    sid=request.args.get("salarie_id",type=int)
    ss=Salarie.query.filter_by(id=sid,tenant_id=t.id).first() if sid else None
    c=Contrat.query.filter_by(salarie_id=sid,tenant_id=t.id,actif=True).first() if sid else None
    acomptes_attente = Acompte.query.filter_by(tenant_id=t.id, salarie_id=sid, statut="EN_ATTENTE").all() if sid else []
    total_acomptes = sum(float(a.montant) for a in acomptes_attente)
    return render_template("tenant/bulletin_saisie.html", salaries=sals, periodes=pers, salarie_sel=ss, contrat=c, tenant=t,
        acomptes_attente=acomptes_attente, total_acomptes=total_acomptes)

@app.route("/bulletins/<int:id>")
@login_required
def bulletin_detail(id):
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    return render_template("tenant/bulletin_detail.html",
        bulletin=BulletinPaie.query.filter_by(id=id,tenant_id=t.id).first_or_404(), tenant=t)

@app.route("/bulletins/<int:id>/valider", methods=["POST"])
@login_required
def bulletin_valider(id):
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if b.statut == "VALIDÉ":
        flash("Ce bulletin est déjà validé.", "info")
        return redirect(url_for("bulletin_detail", id=id))
    # Marquer aussi les acomptes EN_ATTENTE comme DÉDUITS
    acomptes = Acompte.query.filter_by(
        tenant_id=t.id, salarie_id=b.salarie_id,
        mois=b.periode.mois, annee=b.periode.annee,
        statut="EN_ATTENTE").all()
    for a in acomptes:
        a.statut = "DEDUIT"
    b.statut = "VALIDÉ"
    b.date_validation = datetime.utcnow()
    db.session.commit()
    flash("Bulletin validé avec succès.", "success")
    return redirect(url_for("bulletin_detail", id=id))

@app.route("/bulletins/<int:id>/payer", methods=["POST"])
@login_required
def bulletin_paye(id):
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    b.statut = "PAYÉ"
    db.session.commit()
    flash("Bulletin marqué comme payé.", "success")
    return redirect(url_for("bulletin_detail", id=id))

@app.route("/bulletins/<int:id>/supprimer", methods=["POST"])
@login_required
def bulletin_supprimer(id):

    if current_user.is_super_admin:
        b = BulletinPaie.query.get_or_404(id)
        salarie_id = b.salarie_id

        db.session.delete(b)
        db.session.commit()

        flash(
            "Bulletin supprimé (super admin).",
            "success"
        )

        return redirect(
            url_for("salarie_detail", id=salarie_id)
        )

    t = get_tenant()

    if not t:
        return redirect(url_for("login"))

    b = BulletinPaie.query.filter_by(
        id=id,
        tenant_id=t.id
    ).first_or_404()

    # Empêcher suppression si validé ou payé
    if b.statut in ["VALIDÉ", "PAYÉ"]:

        flash(
            "Impossible de supprimer un bulletin validé ou payé.",
            "error"
        )

        return redirect(
            url_for("bulletin_detail", id=id)
        )

    db.session.delete(b)
    db.session.commit()

    flash(
        "Bulletin supprimé.",
        "success"
    )

    return redirect(
        url_for("bulletins")
    )

@app.route("/bulletins/<int:id>/imprimer")
@login_required
def bulletin_imprimer(id):
    if current_user.is_super_admin:
        b = BulletinPaie.query.get_or_404(id)
        t = b.salarie.tenant
    else:
        t = get_tenant()
        if not t: return redirect(url_for("login"))
        b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    return render_template("tenant/bulletin_print.html", bulletin=b, tenant=t)



@app.route("/bulletins/<int:id>/envoyer-email", methods=["POST"])
@login_required
def bulletin_envoyer_email(id):

    import io
    import os
    from xhtml2pdf import pisa
    from flask import current_app

    # Corrige les chemins images/logo pour PDF local + Render
    def link_callback(uri, rel):

        if uri.startswith('/static'):
            path = os.path.join(
                current_app.root_path,
                uri.lstrip('/')
            )

            if os.path.exists(path):
                return path

        return uri

    if current_user.is_super_admin:
        return redirect(url_for("admin_dashboard"))

    t = get_tenant()

    if not t:
        return redirect(url_for("login"))

    b = BulletinPaie.query.filter_by(
        id=id,
        tenant_id=t.id
    ).first_or_404()

    s = b.salarie

    if not s.email:
        flash(
            f"{s.nom_complet} n'a pas d'adresse email.",
            "error"
        )
        return redirect(
            url_for("bulletin_detail", id=id)
        )

    dest_email = request.form.get(
        "email_dest",
        s.email
    ).strip()

    if not app.config.get("MAIL_USERNAME"):
        flash(
            "Email non configuré.",
            "error"
        )
        return redirect(
            url_for("bulletin_detail", id=id)
        )

    try:

        corps = (
            f"Bonjour {s.prenom},\n\n"
            f"Veuillez trouver votre bulletin de paie ci-joint.\n\n"
            f"Période : {b.periode.libelle_complet}\n"
            f"Net à payer : {int(b.net_a_payer or 0):,} FCFA\n\n"
            f"Cordialement,\n"
            f"{t.denomination}"
        ).replace(",", " ")

        msg = Message(
            subject=f"Bulletin de paie {b.periode.libelle_complet}",
            recipients=[dest_email],
            body=corps,
            sender=app.config["MAIL_DEFAULT_SENDER"]
        )

        html = render_template(
            "tenant/bulletin_print.html",
            bulletin=b,
            tenant=t
        )

        pdf_buffer = io.BytesIO()

        result = pisa.CreatePDF(
            html,
            dest=pdf_buffer,
            link_callback=link_callback
        )

        if result.err:
            raise Exception(
                "Erreur génération PDF"
            )

        pdf = pdf_buffer.getvalue()

        msg.attach(
            f"bulletin_{b.id}.pdf",
            "application/pdf",
            pdf
        )

        mail.send(msg)

        flash(
            f"Bulletin envoyé avec succès à {dest_email}",
            "success"
        )

    except Exception as e:

        print(
            "ERREUR COMPLETE :",
            str(e)
        )

        flash(
            f"Erreur envoi : {str(e)}",
            "error"
        )

    return redirect(
        url_for(
            "bulletin_detail",
            id=id
        )
    )

@app.route("/bulletins/envoyer-tous", methods=["POST"])
@login_required
def bulletins_envoyer_tous():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    periode_id = request.form.get("periode_id", type=int)
    if not periode_id: flash("Période manquante.", "error"); return redirect(url_for("bulletins"))
    buls = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=periode_id).all()
    nb_ok=0; nb_err=0
    for b in buls:
        if not b.salarie.email: continue
        try:
            corps = f"Bonjour {b.salarie.prenom},\nBulletin {b.periode.libelle_complet}\nNet: {int(b.net_a_payer or 0):,} FCFA\n{t.denomination}".replace(",", " ")
            msg = Message(subject=f"Bulletin {b.periode.libelle_complet}",
                recipients=[b.salarie.email], body=corps, sender=app.config["MAIL_DEFAULT_SENDER"])
            mail.send(msg); nb_ok+=1
        except: nb_err+=1
    flash(f"{nb_ok} envoyé(s). {nb_err} échec(s).", "success" if nb_ok > 0 else "error")
    return redirect(url_for("bulletins"))

# ── Périodes ──────────────────────────────────────────────────────────────────
@app.route("/periodes")
@login_required
def periodes():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("login"))
    return render_template("tenant/periodes.html", tenant=t,
        periodes=PeriodePaie.query.filter_by(tenant_id=t.id).order_by(PeriodePaie.annee.desc(),PeriodePaie.mois.desc()).all(),
        now=datetime.now())

@app.route("/periodes/nouvelle", methods=["POST"])
@tenant_required
@can_edit
def periode_nouvelle():
    t=get_tenant(); annee=int(request.form["annee"]); mois=int(request.form["mois"])
    noms=PeriodePaie.MOIS_NOMS
    if PeriodePaie.query.filter_by(tenant_id=t.id,annee=annee,mois=mois).first(): flash("Période existante.","warning")
    else:
        db.session.add(PeriodePaie(tenant_id=t.id,annee=annee,mois=mois,libelle_mois=noms[mois],
            trimestre=f"T{(mois-1)//3+1}",statut="OUVERT",date_ouverture=datetime.utcnow()))
        db.session.commit(); flash(f"Période {noms[mois]} {annee} créée.","success")
    return redirect(url_for("periodes"))

@app.route("/periodes/<int:id>/cloturer", methods=["POST"])
@tenant_required
@can_edit
def periode_cloturer(id):
    t=get_tenant(); p=PeriodePaie.query.filter_by(id=id,tenant_id=t.id).first_or_404()
    p.statut="CLÔTURÉ"; p.date_cloture=datetime.utcnow(); db.session.commit()
    flash("Période clôturée.","success"); return redirect(url_for("periodes"))

# ── Paiement abonnement ───────────────────────────────────────────────────────
@app.route("/paiement")
@login_required
def paiement():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    plans = Plan.query.filter_by(actif=True).order_by(Plan.prix_mensuel).all()
    return render_template("tenant/paiement.html", tenant=t, plans=plans)

@app.route("/paiement/confirmer", methods=["POST"])
@login_required
def paiement_confirmer():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    mode = request.form.get("mode", "")
    reference = request.form.get("reference", "").strip()
    duree = int(request.form.get("duree", 1) or 1)
    if not reference: flash("Veuillez indiquer une reference.", "error"); return redirect(url_for("paiement"))
    t.notes = f"PAIEMENT {mode} - Ref: {reference} - {duree} mois - {datetime.now().strftime('%d/%m/%Y')}"
    t.statut = "PAIEMENT_EN_ATTENTE"
    db.session.commit()
    flash(f"Paiement {mode} ref {reference} enregistre. Activation sous 48h.", "success")
    return redirect(url_for("parametres"))

# ── Paramètres ────────────────────────────────────────────────────────────────
@app.route("/parametres")
@tenant_required
def parametres():

    t = Tenant.query.get(get_tenant().id)

    db.session.refresh(t)

    return render_template(
        "tenant/parametres.html",
        tenant=t,
        rubriques=RubriquePaie.query.filter_by(actif=True).all(),
        categories=CategorieEmploi.query.filter_by(tenant_id=t.id).all(),
        users=Utilisateur.query.filter_by(tenant_id=t.id).all()
    )

@app.route("/parametres/logo", methods=["POST"])
@login_required
def parametres_logo():
    import os
    from werkzeug.utils import secure_filename

    if current_user.is_super_admin:
        return redirect(url_for("admin_dashboard"))

    t = get_tenant()
    if not t:
        return redirect(url_for("login"))

    logo_file = request.files.get("logo")

    if logo_file and logo_file.filename:

        dossier = "static/uploads/logos"

        if not os.path.exists(dossier):
            os.makedirs(dossier)

        nom = secure_filename(logo_file.filename)

        chemin = os.path.join(dossier, nom)

        logo_file.save(chemin)

        t.logo_url = "/static/uploads/logos/" + nom

        db.session.add(t)
        db.session.commit()
        db.session.refresh(t)

        

        flash("Logo mis à jour.", "success")

    else:
        flash("Aucun fichier.", "error")

    return redirect(url_for("parametres"))

@app.route("/parametres/logo/supprimer", methods=["POST"])
@login_required
def parametres_logo_supprimer():
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    t.logo_url = None
    db.session.commit()
    flash("Logo supprime.", "success")
    return redirect(url_for("parametres"))

@app.route("/parametres/societe", methods=["POST"])
@tenant_required
@can_edit
def parametres_societe():
    t=get_tenant()
    for f in ["denomination","sigle","activite","nif","numero_cnss","numero_cnamgs","adresse","boite_postale","telephone","ville","region"]:
        setattr(t,f,request.form.get(f,"").strip() or None)
    db.session.commit(); flash("Informations mises à jour.","success")
    return redirect(url_for("parametres"))

@app.route("/parametres/annuler-abonnement", methods=["POST"])
@login_required
def annuler_abonnement():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.is_tenant_admin: abort(403)
    motif = request.form.get("motif", "").strip()
    t.statut = "ANNULATION_DEMANDEE"
    t.notes = f"Annulation demandée le {datetime.now().strftime('%d/%m/%Y')}. Motif: {motif}"
    db.session.commit()
    flash("Demande d annulation enregistrée. L équipe PaieGabon vous contactera sous 48h.", "success")
    return redirect(url_for("parametres"))

# ── Utilisateurs ──────────────────────────────────────────────────────────────
@app.route("/utilisateurs")
@login_required
def utilisateurs():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    liste = Utilisateur.query.filter_by(tenant_id=t.id).order_by(Utilisateur.nom).all()
    return render_template("tenant/utilisateurs.html", tenant=t, utilisateurs=liste, users=liste)

@app.route("/utilisateurs/nouveau", methods=["GET","POST"])
@login_required
def utilisateur_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.is_tenant_admin:
        flash("Réservé à l administrateur.", "error")
        return redirect(url_for("utilisateurs"))
    if request.method == "GET":
        return render_template("tenant/utilisateur_form.html", tenant=t)
    email = request.form.get("email", "").strip().lower()
    nom = request.form.get("nom", "").strip().upper()
    prenom = request.form.get("prenom", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "GESTIONNAIRE")
    if not email or not nom or not password:
        flash("Veuillez remplir tous les champs.", "error")
        return render_template("tenant/utilisateur_form.html", tenant=t)
    if Utilisateur.query.filter_by(email=email).first():
        flash("Email déjà utilisé.", "error")
        return render_template("tenant/utilisateur_form.html", tenant=t)
    u = Utilisateur(nom=nom, prenom=prenom, email=email, role=role, tenant_id=t.id, actif=True)
    u.set_password(password)
    db.session.add(u); db.session.commit()
    flash(f"Utilisateur {u.nom_complet} créé.", "success")
    return redirect(url_for("utilisateurs"))

# ── Journaliers ───────────────────────────────────────────────────────────────
@app.route("/journaliers")
@login_required
def journaliers():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    q = request.args.get("q","")
    query = Journalier.query.filter_by(tenant_id=t.id)
    if q: query = query.filter(db.or_(Journalier.nom.ilike(f"%{q}%"),Journalier.prenom.ilike(f"%{q}%"),Journalier.profession.ilike(f"%{q}%")))
    return render_template("tenant/journaliers.html", tenant=t, journaliers=query.order_by(Journalier.nom).all(), q=q)

@app.route("/journaliers/nouveau", methods=["GET","POST"])
@login_required
def journalier_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    if request.method == "POST":
        j = Journalier(tenant_id=t.id,
            nom=request.form["nom"].strip().upper(),
            prenom=request.form["prenom"].strip(),
            telephone=request.form.get("telephone","").strip(),
            profession=request.form.get("profession","").strip().upper(),
            taux_horaire=float(request.form.get("taux_horaire",0) or 0),
            statut="ACTIF")
        db.session.add(j); db.session.commit()
        flash(f"Journalier {j.nom_complet} créé.", "success")
        return redirect(url_for("journaliers"))
    return render_template("tenant/journalier_form.html", tenant=t, journalier=None)

@app.route("/journaliers/<int:id>/modifier", methods=["GET","POST"])
@login_required
def journalier_modifier(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if request.method == "POST":
        j.nom=request.form["nom"].strip().upper(); j.prenom=request.form["prenom"].strip()
        j.telephone=request.form.get("telephone","").strip()
        j.profession=request.form.get("profession","").strip().upper()
        j.taux_horaire=float(request.form.get("taux_horaire",0) or 0)
        j.statut=request.form.get("statut","ACTIF")
        db.session.commit(); flash("Journalier mis à jour.", "success")
        return redirect(url_for("journaliers"))
    return render_template("tenant/journalier_form.html", tenant=t, journalier=j)

# ── Pointage ──────────────────────────────────────────────────────────────────
@app.route("/pointage")
@login_required
def pointage():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    now = datetime.now()
    date_str = request.args.get("date", now.strftime("%Y-%m-%d"))
    try: date_sel = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: date_sel = now.date()
    salaries_list = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    journaliers_list = Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Journalier.nom).all()
    pts_salaries = {p.salarie_id: p for p in Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_sel).filter(Pointage.salarie_id.isnot(None)).all()}
    pts_journaliers = {p.journalier_id: p for p in Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_sel).filter(Pointage.journalier_id.isnot(None)).all()}
    nb_presents_sal  = sum(1 for p in pts_salaries.values() if p.present)
    nb_presents_jour = sum(1 for p in pts_journaliers.values() if p.present)
    nb_absents       = sum(1 for p in list(pts_salaries.values())+list(pts_journaliers.values()) if p.absent)
    lundi   = date_sel - timedelta(days=date_sel.weekday())
    semaine = [lundi + timedelta(days=i) for i in range(6)]
    return render_template("tenant/pointage.html",
        tenant=t, date_sel=date_sel, semaine=semaine,
        date_hier=(date_sel - timedelta(days=1)).strftime("%Y-%m-%d"),
        date_demain=(date_sel + timedelta(days=1)).strftime("%Y-%m-%d"),
        salaries=salaries_list, journaliers=journaliers_list,
        pts_salaries=pts_salaries, pts_journaliers=pts_journaliers,
        nb_presents_sal=nb_presents_sal, nb_presents_jour=nb_presents_jour,
        nb_absents=nb_absents, now=now)

@app.route("/pointage/sauvegarder", methods=["POST"])
@login_required
def pointage_sauvegarder():
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    date_str = request.form.get("date_pointage")
    try: date_p = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: flash("Date invalide.", "error"); return redirect(url_for("pointage"))
    nb = 0
    for key, val in request.form.items():
        if key.startswith("sal_present_"):
            sid = int(key.replace("sal_present_",""))
            present = val == "1"; absent = not present
            pt = Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_p, salarie_id=sid).first()
            if not pt: pt = Pointage(tenant_id=t.id, date_pointage=date_p, salarie_id=sid); db.session.add(pt)
            pt.present=present; pt.absent=absent
            pt.heures_normales = float(request.form.get(f"sal_heures_{sid}", 8) or 8)
            pt.heures_sup_10   = float(request.form.get(f"sal_sup10_{sid}", 0) or 0)
            pt.heures_sup_30   = float(request.form.get(f"sal_sup30_{sid}", 0) or 0)
            pt.heures_sup_40   = float(request.form.get(f"sal_sup40_{sid}", 0) or 0)
            pt.heures_sup_70   = float(request.form.get(f"sal_sup70_{sid}", 0) or 0)
            pt.motif_absence   = request.form.get(f"sal_motif_{sid}", "") if absent else None
            nb += 1
        if key.startswith("jour_present_"):
            jid = int(key.replace("jour_present_",""))
            present = val == "1"; absent = not present
            pt = Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_p, journalier_id=jid).first()
            if not pt: pt = Pointage(tenant_id=t.id, date_pointage=date_p, journalier_id=jid); db.session.add(pt)
            pt.present=present; pt.absent=absent
            pt.heures_normales = float(request.form.get(f"jour_heures_{jid}", 8) or 8)
            pt.heures_sup      = float(request.form.get(f"jour_sup_{jid}", 0) or 0)
            pt.motif_absence   = request.form.get(f"jour_motif_{jid}", "") if absent else None
            nb += 1
    db.session.commit()
    flash(f"Pointage du {date_p.strftime('%d/%m/%Y')} sauvegardé ({nb} lignes).", "success")
    return redirect(url_for("pointage", date=date_str))

# ── Paie journaliers ──────────────────────────────────────────────────────────
@app.route("/journaliers/paie")
@login_required
def journaliers_paie():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    feuilles = FeuillePaieJournalier.query.filter_by(tenant_id=t.id).order_by(FeuillePaieJournalier.date_fin.desc()).limit(50).all()
    journaliers_list = Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").all()
    return render_template("tenant/journaliers_paie.html", tenant=t, feuilles=feuilles, journaliers=journaliers_list, now=datetime.now())

@app.route("/journaliers/paie/generer", methods=["POST"])
@login_required
def journaliers_paie_generer():
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    date_debut = _parse_date(request.form.get("date_debut"))
    date_fin   = _parse_date(request.form.get("date_fin"))
    if not date_debut or not date_fin: flash("Dates invalides.", "error"); return redirect(url_for("journaliers_paie"))
    nb = 0
    for j in Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").all():
        pts = Pointage.query.filter_by(tenant_id=t.id, journalier_id=j.id)\
              .filter(Pointage.date_pointage>=date_debut, Pointage.date_pointage<=date_fin, Pointage.present==True).all()
        total_h = sum(float(p.heures_normales or 0)+float(p.heures_sup or 0) for p in pts)
        nb_jours = len(pts)
        if nb_jours == 0: continue
        if FeuillePaieJournalier.query.filter_by(tenant_id=t.id, journalier_id=j.id, date_debut=date_debut, date_fin=date_fin).first(): continue
        db.session.add(FeuillePaieJournalier(tenant_id=t.id, journalier_id=j.id,
            date_debut=date_debut, date_fin=date_fin, nb_jours=nb_jours, total_heures=total_h,
            taux_horaire=j.taux_horaire, montant_brut=round(total_h*float(j.taux_horaire),2), statut="EN_ATTENTE"))
        nb += 1
    db.session.commit()
    flash(f"{nb} feuille(s) générée(s).", "success")
    return redirect(url_for("journaliers_paie"))

@app.route("/journaliers/paie/<int:id>/payer", methods=["POST"])
@login_required
def journalier_payer(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    f.statut="PAYÉ"; f.date_paiement=datetime.now().date(); db.session.commit()
    flash(f"Paiement de {f.journalier.nom_complet} enregistré.", "success")
    return redirect(url_for("journaliers_paie"))

@app.route("/journaliers/paie/<int:id>/modifier", methods=["POST"])
@login_required
def journalier_feuille_modifier(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    f.montant_brut = float(request.form.get("montant_brut", f.montant_brut) or f.montant_brut)
    f.observation  = request.form.get("observation", "").strip()
    db.session.commit(); flash("Feuille modifiée.", "success")
    return redirect(url_for("journaliers_paie"))

@app.route("/journaliers/paie/<int:id>/supprimer", methods=["POST"])
@login_required
def journalier_feuille_supprimer(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    db.session.delete(f); db.session.commit()
    flash("Feuille supprimée.", "success")
    return redirect(url_for("journaliers_paie"))

@app.route("/journaliers/paie/payer-selection", methods=["POST"])
@login_required
def journaliers_payer_selection():
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    ids = [int(i) for i in request.form.get("feuille_ids","").split(",") if i.strip().isdigit()]
    nb = 0
    for fid in ids:
        f = FeuillePaieJournalier.query.filter_by(id=fid, tenant_id=t.id, statut="EN_ATTENTE").first()
        if f: f.statut="PAYÉ"; f.date_paiement=datetime.now().date(); nb+=1
    db.session.commit()
    flash(f"{nb} journalier(s) payé(s).", "success")
    return redirect(url_for("journaliers_paie"))

# ── Acomptes ──────────────────────────────────────────────────────────────────
@app.route("/acomptes")
@login_required
def acomptes():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    now = datetime.now()
    mois = request.args.get("mois", now.month, type=int)
    annee = request.args.get("annee", now.year, type=int)
    salarie_id = request.args.get("salarie_id", type=int)
    try:
        query = Acompte.query.filter_by(tenant_id=t.id, annee=annee, mois=mois)
        if salarie_id: query = query.filter_by(salarie_id=salarie_id)
        liste = query.order_by(Acompte.date_acompte.desc()).all()
    except Exception:
        db.create_all(); db.session.rollback(); liste = []
    total_mois       = sum(float(a.montant) for a in liste if a.statut != "ANNULE")
    total_en_attente = sum(float(a.montant) for a in liste if a.statut == "EN_ATTENTE")
    total_deduit     = sum(float(a.montant) for a in liste if a.statut == "DEDUIT")
    salaries_list = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    return render_template("tenant/acomptes.html", tenant=t, liste=liste, salaries=salaries_list,
        mois=mois, annee=annee, now=now, total_mois=total_mois,
        total_en_attente=total_en_attente, total_deduit=total_deduit, MOIS_NOMS=PeriodePaie.MOIS_NOMS)

@app.route("/acomptes/nouveau", methods=["GET","POST"])
@login_required
def acompte_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    if not current_user.can_edit: abort(403)
    salaries_list = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    if request.method == "POST":
        salarie_id = request.form.get("salarie_id", type=int)
        montant    = float(request.form.get("montant", 0) or 0)
        date_ac    = _parse_date(request.form.get("date_acompte"))
        mois       = request.form.get("mois", type=int)
        annee      = request.form.get("annee", type=int)
        motif      = request.form.get("motif", "").strip()
        if not salarie_id or montant <= 0 or not date_ac:
            flash("Veuillez remplir tous les champs.", "error")
        else:
            contrat = Contrat.query.filter_by(salarie_id=salarie_id, tenant_id=t.id, actif=True).first()
            if contrat and montant > float(contrat.salaire_base) * 0.5:
                flash(f"Acompte maximum 50% du salaire de base ({float(contrat.salaire_base)*0.5:,.0f} FCFA).".replace(",", " "), "error")
                return render_template("tenant/acompte_form.html", tenant=t, salaries=salaries_list, now=datetime.now())
            db.session.add(Acompte(tenant_id=t.id, salarie_id=salarie_id, montant=montant,
                date_acompte=date_ac, mois=mois, annee=annee, motif=motif, statut="EN_ATTENTE"))
            db.session.commit()
            flash(f"Acompte de {montant:,.0f} FCFA enregistré.".replace(",", " "), "success")
            return redirect(url_for("acomptes", mois=mois, annee=annee))
    return render_template("tenant/acompte_form.html", tenant=t, salaries=salaries_list, now=datetime.now())

@app.route("/acomptes/<int:id>/valider", methods=["POST"])
@login_required
def acompte_valider(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    a = Acompte.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    a.statut = "DEDUIT"; db.session.commit()
    flash("Acompte marqué comme déduit.", "success")
    return redirect(url_for("acomptes", mois=a.mois, annee=a.annee))

@app.route("/acomptes/<int:id>/annuler", methods=["POST"])
@login_required
def acompte_annuler(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    a = Acompte.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    a.statut = "ANNULE"; db.session.commit()
    flash("Acompte annulé.", "success")
    return redirect(url_for("acomptes", mois=a.mois, annee=a.annee))

@app.route("/acomptes/<int:id>/supprimer", methods=["POST"])
@login_required
def acompte_supprimer(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    a = Acompte.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    mois, annee = a.mois, a.annee
    db.session.delete(a); db.session.commit()
    flash("Acompte supprimé.", "success")
    return redirect(url_for("acomptes", mois=mois, annee=annee))

# ── Congés ────────────────────────────────────────────────────────────────────
@app.route("/conges")
@login_required
def conges():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    now = datetime.now()
    annee = request.args.get("annee", now.year, type=int)
    q = request.args.get("q", "")
    salaries_list = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    soldes = []
    for s in salaries_list:
        if q and q.lower() not in f"{s.nom} {s.prenom} {s.matricule}".lower(): continue
        conge = Conge.query.filter_by(tenant_id=t.id, salarie_id=s.id, annee=annee).first()
        mois_anc = max(1,(datetime.now().date()-s.date_embauche).days//30) if s.date_embauche else 12
        jours_auto = round(min(mois_anc,12)*2.0,1)
        soldes.append({"salarie":s,"conge":conge,
            "jours_acquis":float(conge.jours_acquis) if conge else jours_auto,
            "jours_pris":float(conge.jours_pris) if conge else 0,
            "jours_restants":(float(conge.jours_acquis)-float(conge.jours_pris)) if conge else jours_auto})
    demandes = Conge.query.filter_by(tenant_id=t.id).filter(Conge.statut.in_(["DEMANDÉ","APPROUVÉ"])).order_by(Conge.date_depart).all()
    return render_template("tenant/conges.html", tenant=t, soldes=soldes, demandes=demandes,
        annee=annee, now=now, q=q, salaries=salaries_list)

@app.route("/conges/nouveau", methods=["GET","POST"])
@login_required
def conge_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    salaries_list = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    if request.method == "POST":
        salarie_id = request.form.get("salarie_id", type=int)
        annee = request.form.get("annee", datetime.now().year, type=int)
        date_dep = _parse_date(request.form.get("date_depart"))
        date_ret = _parse_date(request.form.get("date_retour"))
        type_c = request.form.get("type_conge", "ANNUEL")
        jours = (date_ret - date_dep).days + 1 if date_dep and date_ret else 0
        conge = Conge.query.filter_by(tenant_id=t.id, salarie_id=salarie_id, annee=annee).first()
        if not conge:
            s = Salarie.query.get(salarie_id)
            mois = max(1,(datetime.now().date()-s.date_embauche).days//30) if s.date_embauche else 12
            conge = Conge(tenant_id=t.id, salarie_id=salarie_id, annee=annee,
                jours_acquis=round(min(mois,12)*2.0,1), jours_pris=0, type_conge=type_c, statut="DEMANDÉ")
            db.session.add(conge)
        conge.date_depart=date_dep; conge.date_retour=date_ret; conge.type_conge=type_c; conge.statut="DEMANDÉ"
        db.session.commit()
        flash(f"Demande de congé enregistrée ({jours} jours).", "success")
        return redirect(url_for("conges"))
    return render_template("tenant/conge_form.html", tenant=t, salaries=salaries_list, now=datetime.now())

@app.route("/conges/<int:id>/approuver", methods=["POST"])
@login_required
def conge_approuver(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    c = Conge.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if c.date_depart and c.date_retour:
        c.jours_pris = float(c.jours_pris or 0) + (c.date_retour-c.date_depart).days + 1
    c.statut = "APPROUVÉ"; db.session.commit()
    flash(f"Congé de {c.salarie.nom_complet} approuvé.", "success")
    return redirect(url_for("conges"))

@app.route("/conges/<int:id>/refuser", methods=["POST"])
@login_required
def conge_refuser(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    c = Conge.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    c.statut="REFUSÉ"; db.session.commit()
    flash(f"Congé refusé.", "success")
    return redirect(url_for("conges"))

@app.route("/conges/<int:id>/supprimer", methods=["POST"])
@login_required
def conge_supprimer(id):
    t = get_tenant()
    if not t: return redirect(url_for("login"))
    c = Conge.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    db.session.delete(c); db.session.commit()
    flash("Demande supprimée.", "success")
    return redirect(url_for("conges"))

# ── Export & API ──────────────────────────────────────────────────────────────
@app.route("/bulletins/export/<int:periode_id>")
@tenant_required
def export_journal(periode_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    t=get_tenant()
    p=PeriodePaie.query.filter_by(id=periode_id,tenant_id=t.id).first_or_404()
    buls=BulletinPaie.query.filter_by(periode_id=periode_id,tenant_id=t.id).join(Salarie).order_by(Salarie.nom).all()
    wb=Workbook(); ws=wb.active; ws.title=f"Journal {p.libelle_complet}"
    ws.merge_cells("A1:R1"); ws["A1"]=f"JOURNAL DE PAIE — {p.libelle_complet} — {t.denomination}"
    ws["A1"].font=Font(bold=True,size=13); ws["A1"].alignment=Alignment(horizontal="center")
    hdrs=["Matricule","Nom","Prénom","Emploi","Cat.","Base","Brut","CNSS Sal.","CNAMGS Sal.","TCS","IRPP","Net","Net à Payer","CNSS Pat.","CNAMGS Pat.","FNH","CFP","Statut"]
    for col,h in enumerate(hdrs,1):
        c=ws.cell(row=3,column=col,value=h); c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1a2332"); c.alignment=Alignment(horizontal="center")
    for row,b in enumerate(buls,4):
        s=b.salarie
        vals=[s.matricule,s.nom,s.prenom,s.emploi,s.categorie.code if s.categorie else "",
              float(b.salaire_base or 0),float(b.salaire_brut or 0),float(b.cnss_salarie or 0),
              float(b.cnamgs_salarie or 0),float(b.tcs or 0),float(b.irpp or 0),
              float(b.salaire_net or 0),float(b.net_a_payer or 0),float(b.cnss_patronale or 0),
              float(b.cnamgs_patronale or 0),float(b.fnh or 0),float(b.cfp or 0),b.statut]
        for col,v in enumerate(vals,1):
            cell=ws.cell(row=row,column=col,value=v)
            if isinstance(v,float): cell.number_format='#,##0'
            if row%2==0: cell.fill=PatternFill("solid",fgColor="F5F5F5")
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,download_name=f"Journal_{p.libelle_mois}_{p.annee}_{t.slug}.xlsx")

@app.route("/api/calculer-bulletin", methods=["POST"])
@login_required
def api_calculer():
    try:
        t = get_tenant()
        data = request.get_json() or {}
        sid = data.pop("salarie_id", None)
        nb_parts = 1.0
        if sid and t:
            s = Salarie.query.filter_by(id=sid, tenant_id=t.id).first()
            if s: nb_parts = float(s.nombre_parts or 1)
        # Récupérer les acomptes EN_ATTENTE pour ce salarié (calcul temps réel)
        mois  = data.pop("mois_periode", None)
        annee = data.pop("annee_periode", None)
        total_acomptes = 0.0
        if sid and t and mois and annee:
            total_acomptes = float(db.session.query(db.func.sum(Acompte.montant))                .filter_by(tenant_id=t.id, salarie_id=int(sid), mois=int(mois),
                           annee=int(annee), statut="EN_ATTENTE").scalar() or 0)
        if total_acomptes > 0:
            data["acompte"] = max(float(data.get("acompte", 0)), total_acomptes)
        res = calculer_bulletin(data, nb_parts=nb_parts)
        res["acompte_auto"] = total_acomptes
        return jsonify(res)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/salarie/<int:id>/contrat")
@login_required
def api_contrat(id):
    t=get_tenant()
    s=Salarie.query.filter_by(id=id,tenant_id=t.id).first()
    if not s: return jsonify({})
    c=Contrat.query.filter_by(salarie_id=id,tenant_id=t.id,actif=True).first()
    base={"nom":s.nom_complet,"poste":s.emploi,"matricule":s.matricule,"nombre_parts":float(s.nombre_parts or 1)}
    if c: base["salaire_base"]=float(c.salaire_base); base["poste"]=c.poste or s.emploi
    return jsonify(base)

@app.route("/api/salarie/<int:id>/acomptes-mois")
@login_required
def api_acomptes_mois(id):
    t = get_tenant()
    mois=request.args.get("mois",type=int); annee=request.args.get("annee",type=int)
    if not t or not mois or not annee: return jsonify({"total":0})
    total = db.session.query(db.func.sum(Acompte.montant))\
            .filter_by(tenant_id=t.id,salarie_id=id,mois=mois,annee=annee,statut="EN_ATTENTE").scalar() or 0
    return jsonify({"total":float(total)})

@app.route("/api/pointage/semaine")
@login_required
def api_pointage_semaine():
    t = get_tenant()
    if not t: return jsonify({})
    date_str = request.args.get("date")
    try: date_sel = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: date_sel = datetime.now().date()
    lundi=date_sel-timedelta(days=date_sel.weekday()); samedi=lundi+timedelta(days=5)
    pts = Pointage.query.filter_by(tenant_id=t.id).filter(Pointage.date_pointage>=lundi,Pointage.date_pointage<=samedi).all()
    stats={}
    for p in pts:
        key=str(p.date_pointage)
        if key not in stats: stats[key]={"presents":0,"absents":0,"heures":0}
        if p.present: stats[key]["presents"]+=1; stats[key]["heures"]+=p.total_heures
        else: stats[key]["absents"]+=1
    return jsonify(stats)

# ── Helpers ───────────────────────────────────────────────────────────────────
def _pd(v):
    if not v: return None
    try: return datetime.strptime(v,"%Y-%m-%d").date()
    except: return None

@app.template_filter("fcfa")
def fcfa_filter(v):
    try: return f"{int(float(v)):,}".replace(","," ")+" FCFA"
    except: return "— FCFA"

@app.template_filter("date_fr")
def date_fr_filter(v):
    if not v: return "—"
    if isinstance(v,str):
        try: v=datetime.strptime(v[:10],"%Y-%m-%d").date()
        except: return v
    return v.strftime("%d/%m/%Y")

@app.context_processor
def inject_globals(): return {"now":datetime.now(),"enumerate":enumerate}

@app.errorhandler(403)
def forbidden(e): return render_template("auth/403.html"),403

# ── Init DB ───────────────────────────────────────────────────────────────────
def init_db():
    db.create_all()
    if not Plan.query.first():
        for code,nom,prix,ms,mu,desc in [
            ("STARTER","Starter",15000,10,1,"10 salariés max, 1 utilisateur"),
            ("PRO","Pro",35000,50,3,"50 salariés max, 3 utilisateurs"),
            ("CABINET","Cabinet",100000,None,10,"Illimité, 10 utilisateurs"),
        ]: db.session.add(Plan(code=code,nom=nom,prix_mensuel=prix,max_salaries=ms,max_utilisateurs=mu,description=desc))
    if not RubriquePaie.query.first():
        for code,lib,typ,ts,tp,plaf in [
            ("CNSS","Caisse Nationale Sécurité Sociale","COTISATION",0.05,0.18,1500000),
            ("CNAMGS","Assurance Maladie Garantie Sociale","COTISATION",0.02,0.041,2500000),
            ("TCS","Taxe Complémentaire Salaires","RETENUE",0.05,None,None),
            ("FNH","Fonds National Habitat","COTISATION",None,0.03,1500000),
            ("CFP","Contribution Formation Professionnelle","COTISATION",None,0.005,None),
        ]: db.session.add(RubriquePaie(code=code,libelle=lib,type=typ,taux_salarie=ts,taux_patronal=tp,plafond_mensuel=plaf))
    if not Utilisateur.query.filter_by(role="SUPER_ADMIN").first():
        sa=Utilisateur(nom="ADMIN",prenom="Super",email="superadmin@paiegalon.com",role="SUPER_ADMIN",actif=True)
        sa.set_password("Admin2026!"); db.session.add(sa)
    if not Tenant.query.first():
        db.session.flush()
        plan=Plan.query.filter_by(code="PRO").first()
        t=Tenant(slug="demo",denomination="ENTREPRISE DEMO",sigle="DEMO",activite="À RENSEIGNER",
                 nif="",ville="Libreville",pays="Gabon",
                 plan_id=plan.id if plan else None,statut="ACTIF")
        t.token_api=sec.token_hex(32)
        db.session.add(t); db.session.flush()
        for code,lib in [("C1","Ouvriers"),("C2","Techniciens"),("C3","Conducteurs"),("C4","Cadres")]:
            db.session.add(CategorieEmploi(tenant_id=t.id,code=code,libelle=lib))
        u=Utilisateur(nom="DEMO",prenom="Responsable",email="demo@paiegalon.ga",role="TENANT_ADMIN",tenant_id=t.id,actif=True)
        u.set_password("Demo2026!"); db.session.add(u)
    db.session.commit()
    print("Base initialisée.\n  Super-admin: superadmin@paiegalon.com / Admin2026!\n  Compte démo: demo@paiegalon.ga / Demo2026!")

with app.app_context():
    try:
        db.create_all()
        # Migrations colonnes manquantes
        for col in ["heures_sup_10","heures_sup_30","heures_sup_40","heures_sup_70"]:
            try:
                db.session.execute(db.text(f"ALTER TABLE pointages ADD COLUMN IF NOT EXISTS {col} NUMERIC(5,2) DEFAULT 0"))
                db.session.commit()
            except Exception: db.session.rollback()
        try:
            db.session.execute(db.text("ALTER TABLE tenants ADD COLUMN IF NOT EXISTS logo_url VARCHAR(500)"))
            db.session.commit()
        except Exception: db.session.rollback()
        # Migration: ajouter email dans salaries
        try:
            db.session.execute(db.text("ALTER TABLE salaries ADD COLUMN IF NOT EXISTS email VARCHAR(200)"))
            db.session.commit()
        except Exception: db.session.rollback()
        init_db()
        print("✅ Tables créées et base initialisée.")
    except Exception as e:
        print(f"Erreur init: {e}")
        try: db.session.rollback(); db.create_all()
        except Exception as e2: print(f"Erreur create_all: {e2}")


        

if __name__=="__main__":
    with app.app_context(): init_db()
    app.run(debug=True,port=5000)
